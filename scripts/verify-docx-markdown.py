from __future__ import annotations

import re
import sys
from pathlib import Path
from docx import Document


def normalize(value: str) -> list[str]:
    return re.findall(r"[a-z0-9À-ÿ]+", value.lower())


source_dir, markdown_dir = map(Path, sys.argv[1:3])
failed = False
for source in sorted(source_dir.glob("*.docx")):
    name = source.name[:-5]
    if name.lower().endswith(".md"):
        name = name[:-3]
    target = markdown_dir / f"{name}.md"
    document = Document(source)
    source_text = "\n".join([paragraph.text for paragraph in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    markdown = target.read_text(encoding="utf-8") if target.exists() else ""
    source_tokens = normalize(source_text)
    markdown_tokens = normalize(markdown)
    coverage = len(markdown_tokens) / max(1, len(source_tokens))
    ok = target.exists() and len(markdown) > 100 and coverage >= 0.95
    failed = failed or not ok
    print(f"{source.name}: {'OK' if ok else 'FALHA'} | cobertura={coverage:.2f} | docx_tokens={len(source_tokens)} | md_tokens={len(markdown_tokens)}")
if failed:
    raise SystemExit(1)
